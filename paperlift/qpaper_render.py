"""Exporters: compact Markdown, HTML (preview), PDF (PyMuPDF Story), DOCX, JSON."""
from __future__ import annotations

import html
import json
import os
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from typing import List, Optional

from .qpaper_engine import Paper, Section, Question, Group, guess_title_fields, subject_display, runs_plain

LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
PX = 4 / 3  # pt -> css px


# --------------------------------------------------------------------------
# Options + document model
# --------------------------------------------------------------------------
DEFAULT_OPTIONS = {
    "sections": "auto",       # auto | all | comma separated section indexes
    "split": True,            # one output per subject section
    "renumber": True,         # 1..N instead of original question numbers
    "marks": True,            # show [n marks]
    "qtype_hint": True,       # "(select all that apply)" / "(numeric answer)"
    "title": "",              # full title override
    "code": "", "year": "", "term": "", "exam": "", "subject": "",
}


@dataclass
class DocModel:
    title: str
    slug: str
    items: list = field(default_factory=list)    # ("passage", Group, (a, b)) | ("q", Question, n)
    key: list = field(default_factory=list)      # (n, answer_str, marks, qtype)
    section: Optional[Section] = None


def pick_sections(paper: Paper, mode) -> List[Section]:
    secs = paper.sections
    if isinstance(mode, str) and mode not in ("auto", "all", ""):
        idx = {int(x) for x in re.findall(r"\d+", mode)}
        return [s for i, s in enumerate(secs) if i in idx]
    if mode == "all":
        return [s for s in secs if s.questions]
    subj = [s for s in secs if s.subject and s.questions]
    if subj:
        return subj
    full = [s for s in secs if not s.partial and s.questions]
    return full or [s for s in secs if s.questions]


def answer_str(q: Question) -> str:
    if q.options:
        letters = [LETTERS[i] for i, o in enumerate(q.options) if o.correct]
        return ", ".join(letters) if letters else "—"
    if q.answers:
        return " / ".join(a.replace(" to ", " to ") for a in q.answers)
    return "—"


def build_title(paper: Paper, sec: Section, opts: dict) -> str:
    if opts.get("title"):
        return opts["title"]
    g = guess_title_fields(paper.source, paper)
    for k in ("code", "year", "term", "exam", "subject"):
        if opts.get(k):
            g[k] = opts[k]
    subject = g["subject"] or subject_display(sec)
    head = subject + (f" ({g['code']})" if g["code"] else "")
    when = " ".join(x for x in (g["year"], g["term"]) if x)
    parts = [head] + [x for x in (when, g["exam"]) if x]
    return " · ".join(parts)


def build_models(paper: Paper, opts: dict) -> List[DocModel]:
    o = dict(DEFAULT_OPTIONS, **{k: v for k, v in (opts or {}).items() if v is not None})
    secs = pick_sections(paper, o["sections"])
    stem = os.path.splitext(os.path.basename(paper.source))[0]
    groups_of = [secs] if not o["split"] or len(secs) <= 1 else [[s] for s in secs]
    models = []
    for gi, group_secs in enumerate(groups_of):
        title = build_title(paper, group_secs[0], o)
        slug = stem if len(groups_of) == 1 else f"{stem}_{re.sub(r'[^A-Za-z0-9]+', '_', subject_display(group_secs[0])).strip('_')}"
        m = DocModel(title=title, slug=slug, section=group_secs[0])
        n = 0
        for sec in group_secs:
            if len(group_secs) > 1:
                m.items.append(("heading", subject_display(sec), None))
            qs = [q for q in sec.questions if not q.subject_check]
            shown = set()
            numbers = {}
            for q in qs:
                n += 1
                numbers[q.qid] = n if o["renumber"] else q.number
            for q in qs:
                disp = numbers[q.qid]
                if q.group and q.group not in shown and q.group in sec.groups:
                    shown.add(q.group)
                    members = [numbers[x.qid] for x in qs if x.group == q.group]
                    m.items.append(("passage", sec.groups[q.group], (min(members), max(members))))
                m.items.append(("q", q, disp))
                m.key.append((disp, answer_str(q), q.marks, q.qtype))
        models.append(m)
    return models


def qtype_hint(q: Question) -> str:
    if q.qtype == "MSQ":
        return "select all that apply"
    if q.qtype == "SA":
        return "numeric answer" if q.answers and all(re.match(r"^-?[\d.]+(\s+to\s+-?[\d.]+)?$", a) for a in q.answers) else "short answer"
    return ""


def fmt_marks(v) -> str:
    if v is None:
        return ""
    v = int(v) if float(v).is_integer() else v
    return f"{v} mark" + ("" if v == 1 else "s")


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------
def _md_escape(t: str) -> str:
    t = t.replace("\\", "\\\\")
    t = re.sub(r"([*_`\[\]])", r"\\\1", t)
    t = t.replace("<", "&lt;").replace(">", "&gt;")
    return t


def _wrap_keep_space(t: str, a: str, b: str) -> str:
    """Wrap text in markers, keeping surrounding spaces outside (`**x** ` not `**x **`)."""
    core = t.strip()
    if not core:
        return t
    lead, trail = t[:len(t) - len(t.lstrip())], t[len(t.rstrip()):]
    return f"{lead}{a}{core}{b}{trail}"


def _md_styled(r: dict) -> str:
    if r.get("code"):
        return _wrap_keep_space(r["t"].replace("`", "'"), "`", "`")
    t = _md_escape(r["t"])
    if r.get("s"):
        t = f"<{r['s']}>{t}</{r['s']}>"
    if r.get("b") and r.get("i"):
        t = _wrap_keep_space(t, "***", "***")
    elif r.get("b"):
        t = _wrap_keep_space(t, "**", "**")
    elif r.get("i"):
        t = _wrap_keep_space(t, "*", "*")
    return t


def _html_styled(r: dict) -> str:
    t = html.escape(r["t"])
    if r.get("s"):
        t = f"<{r['s']}>{t}</{r['s']}>"
    if r.get("code"):
        t = f"<code>{t}</code>"
    if r.get("i"):
        t = f"<i>{t}</i>"
    if r.get("b"):
        t = f"<b>{t}</b>"
    return t


def md_runs(runs, img_prefix="images/") -> str:
    out = []
    for r in runs:
        if "p" in r:
            out.append("\n\n")
        elif "t" in r:
            out.append(_md_styled(r))
        elif "img" in r:
            out.append(f'<img src="{img_prefix}{r["img"]}" height="{round(r["h"] * PX)}" alt="" style="vertical-align:middle">')
        elif "block" in r:
            out.append(f'\n\n<img src="{img_prefix}{r["block"]}" width="{round(r["w"] * PX)}" alt="figure">\n\n')
    s = "".join(out)
    s = re.sub(r"\n{3,}", "\n\n", s).strip()
    s = re.sub(r"^(\s*)(\d+)\.(?=\s)", r"\1\2\\.", s, flags=re.M)     # avoid accidental lists
    s = re.sub(r"^(\s*)([#>+\-])(?=\s)", r"\1\\\2", s, flags=re.M)
    return s


def to_markdown(m: DocModel, opts: dict) -> str:
    o = dict(DEFAULT_OPTIONS, **(opts or {}))
    L = [f"# {m.title}", ""]
    for kind, obj, extra in m.items:
        if kind == "heading":
            L += [f"## {obj}", ""]
        elif kind == "passage":
            a, b = extra
            rng = f"Questions {a}–{b}" if a != b else f"Question {a}"
            L += ["---", "", f"**{rng} are based on the following:**", "", md_runs(obj.passage), ""]
        else:
            q, n = obj, extra
            tail = []
            if o["qtype_hint"] and qtype_hint(q):
                tail.append(f"*({qtype_hint(q)})*")
            if o["marks"] and q.marks is not None:
                tail.append(f"`[{fmt_marks(q.marks)}]`")
            body = md_runs(q.body)
            lead = f"**Q{n}.**" + ((" " + " ".join(tail)) if tail else "")
            L.append(f"{lead} {body}" if not body.startswith("<img") or "height=" in body.split(">")[0] else f"{lead}\n\n{body}")
            L.append("")
            for i, op in enumerate(q.options):
                L.append(f"- **({LETTERS[i]})** {md_runs(op.content).replace(chr(10) + chr(10), '<br>')}")
            if q.options:
                L.append("")
    L += ["---", "", "## Answer Key", "", "| Q | Answer |" + (" Marks |" if o["marks"] else ""),
          "|:-:|:--|" + (":-:|" if o["marks"] else "")]
    for n, ans, marks, qt in m.key:
        L.append(f"| {n} | {_md_escape(ans)} |" + (f" {fmt_marks(marks).split()[0] if marks is not None else ''} |" if o["marks"] else ""))
    return "\n".join(L).rstrip() + "\n"


# --------------------------------------------------------------------------
# HTML (preview + PDF source)
# --------------------------------------------------------------------------
def html_runs(runs, img_prefix="images/", for_pdf=False) -> str:
    out, para = [], []

    def flush():
        if para:
            out.append("<p>" + "".join(para) + "</p>")
            para.clear()
    for r in runs:
        if "p" in r:
            flush()
        elif "t" in r:
            para.append(_html_styled(r))
        elif "img" in r:
            if for_pdf:
                para.append(f'<img src="{r["img"]}" width="{r["w"]:.1f}" height="{r["h"]:.1f}"/>')
            else:
                para.append(f'<img class="inl" src="{img_prefix}{r["img"]}" style="height:{r["h"] * PX:.0f}px" alt="">')
        elif "block" in r:
            flush()
            if for_pdf:
                w = min(r["w"], 480.0)
                h = r["h"] * w / r["w"]
                out.append(f'<p class="fig"><img src="{r["block"]}" width="{w:.1f}" height="{h:.1f}"/></p>')
            else:
                out.append(f'<p class="fig"><img src="{img_prefix}{r["block"]}" style="width:{r["w"] * PX:.0f}px;max-width:100%" alt="figure"></p>')
    flush()
    return "".join(out)


PREVIEW_CSS = """
:root{color-scheme:light}
body{font:15px/1.55 "Charter","Georgia",serif;max-width:820px;margin:28px auto;padding:0 20px;color:#1d1d1f;background:#fff}
h1{font:600 22px/1.3 -apple-system,"Segoe UI",sans-serif;border-bottom:2px solid #222;padding-bottom:8px}
h2{font:600 17px -apple-system,"Segoe UI",sans-serif;margin-top:28px}
.q{margin:18px 0;break-inside:avoid}
.q>.stem p{margin:.25em 0}
.qn{font-weight:700}
.tag{font:12px -apple-system,sans-serif;color:#666;margin-left:6px}
ol.opts{list-style:none;padding-left:22px;margin:6px 0}
ol.opts li{margin:3px 0}
ol.opts li p{display:inline;margin:0}
.ol{font-weight:600;margin-right:6px}
.passage{border-left:3px solid #999;background:#f6f6f4;padding:8px 14px;margin:22px 0 8px}
.passage .ph{font:600 13px -apple-system,sans-serif;color:#444}
img.inl{vertical-align:middle}
.fig{margin:8px 0}
table.key{border-collapse:collapse;font:14px -apple-system,sans-serif}
table.key td,table.key th{border:1px solid #ccc;padding:4px 12px;text-align:left}
table.key th{background:#f0f0f0}
"""

PDF_CSS = """
body{font-family:serif;font-size:11pt;line-height:1.4}
h1{font-family:sans-serif;font-size:16pt;font-weight:bold;margin-bottom:8pt}
h2{font-family:sans-serif;font-size:13pt;font-weight:bold;margin-top:12pt}
p{margin-top:2pt;margin-bottom:2pt}
.qn{font-weight:bold}
.tag{font-family:sans-serif;font-size:9pt;color:#666666}
.opt{margin-left:18pt;margin-top:1pt;margin-bottom:1pt}
.ol{font-weight:bold}
.passage{margin-top:10pt;margin-bottom:4pt;padding-left:8pt;border-left:2px solid #999999}
.ph{font-family:sans-serif;font-size:10pt;font-weight:bold;color:#444444}
.q{margin-top:9pt}
table{border-collapse:collapse}
td,th{border:1px solid #999999;padding:2pt 8pt;font-family:sans-serif;font-size:10pt}
"""


def to_html(m: DocModel, opts: dict, for_pdf=False, img_prefix="images/") -> str:
    o = dict(DEFAULT_OPTIONS, **(opts or {}))
    H = []
    H.append(f"<h1>{html.escape(m.title)}</h1>")
    for kind, obj, extra in m.items:
        if kind == "heading":
            H.append(f"<h2>{html.escape(obj)}</h2>")
        elif kind == "passage":
            a, b = extra
            rng = f"Questions {a}–{b}" if a != b else f"Question {a}"
            H.append(f'<div class="passage"><p class="ph">{rng} are based on the following:</p>'
                     f'{html_runs(obj.passage, img_prefix, for_pdf)}</div>')
        else:
            q, n = obj, extra
            tags = []
            if o["qtype_hint"] and qtype_hint(q):
                tags.append(qtype_hint(q))
            if o["marks"] and q.marks is not None:
                tags.append(fmt_marks(q.marks))
            tag = f'<span class="tag">[{" · ".join(tags)}]</span>' if tags else ""
            stem = html_runs(q.body, img_prefix, for_pdf) or "<p></p>"
            lead = f'<span class="qn">Q{n}.</span> {tag} '
            if stem.startswith('<p class="fig">'):
                stem = stem.replace('<p class="fig">', f'<p class="fig">{lead}<br/>', 1)
            elif stem.startswith("<p>"):
                stem = stem.replace("<p>", "<p>" + lead, 1)
            else:
                stem = f"<p>{lead}</p>" + stem
            if for_pdf:
                opts_html = "".join(
                    f'<p class="opt"><span class="ol">({LETTERS[i]})</span> '
                    + re.sub(r"</?p( class=\"fig\")?>", " ", html_runs(op.content, img_prefix, True)) + "</p>"
                    for i, op in enumerate(q.options))
            else:
                opts_html = ('<ol class="opts">' + "".join(
                    f'<li><span class="ol">({LETTERS[i]})</span>{html_runs(op.content, img_prefix)}</li>'
                    for i, op in enumerate(q.options)) + "</ol>") if q.options else ""
            H.append(f'<div class="q"><div class="stem">{stem}</div>{opts_html}</div>')
    rows = "".join(
        f"<tr><td>{n}</td><td>{html.escape(a)}</td>" + (f"<td>{fmt_marks(mk).split()[0] if mk is not None else ''}</td>" if o["marks"] else "") + "</tr>"
        for n, a, mk, _ in m.key)
    H.append('<h2>Answer Key</h2><table class="key"><tr><th>Q</th><th>Answer</th>'
             + ("<th>Marks</th>" if o["marks"] else "") + f"</tr>{rows}</table>")
    body = "\n".join(H)
    if for_pdf:
        return body
    return (f'<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f"<title>{html.escape(m.title)}</title><style>{PREVIEW_CSS}</style></head><body>{body}</body></html>")


def to_pdf(m: DocModel, opts: dict, img_dir: str, out_path: str):
    import pymupdf as fitz
    body = to_html(m, opts, for_pdf=True)
    archive = fitz.Archive(img_dir) if os.path.isdir(img_dir) else None
    story = fitz.Story(html=body, user_css=PDF_CSS, archive=archive)
    writer = fitz.DocumentWriter(out_path)
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (54, 54, -54, -54)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    # page numbers
    doc = fitz.open(out_path)
    for i, page in enumerate(doc):
        page.insert_text((mediabox.width / 2 - 12, mediabox.height - 28), f"{i + 1} / {doc.page_count}",
                         fontsize=8, color=(0.45, 0.45, 0.45))
    doc.saveIncr()
    doc.close()


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def to_docx(m: DocModel, opts: dict, img_dir: str, out_path: str):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    o = dict(DEFAULT_OPTIONS, **(opts or {}))
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.paragraph_format.space_after = Pt(2)
    max_w = 17.0 * 28.35  # usable width in pt

    def add_runs(p, runs, new_para):
        for r in runs:
            if "p" in r:
                p = new_para()
            elif "t" in r:
                run = p.add_run(r["t"])
                if r.get("s") == "sub":
                    run.font.subscript = True
                elif r.get("s") == "sup":
                    run.font.superscript = True
                if r.get("b"):
                    run.bold = True
                if r.get("i"):
                    run.italic = True
                if r.get("code"):
                    run.font.name = "Consolas"
            elif "img" in r or "block" in r:
                name = r.get("img") or r.get("block")
                path = os.path.join(img_dir, name)
                if not os.path.exists(path):
                    continue
                w = min(r["w"], max_w)
                h = r["h"] * w / r["w"]
                if "block" in r:
                    p = new_para()
                    p.add_run().add_picture(path, height=Pt(h))
                    p = new_para()
                else:
                    p.add_run().add_picture(path, height=Pt(h))
        return p

    t = doc.add_heading(m.title, level=1)
    for kind, obj, extra in m.items:
        if kind == "heading":
            doc.add_heading(obj, level=2)
        elif kind == "passage":
            a, b = extra
            rng = f"Questions {a}–{b}" if a != b else f"Question {a}"
            hp = doc.add_paragraph()
            hr = hp.add_run(f"{rng} are based on the following:")
            hr.bold = True
            hp.paragraph_format.space_before = Pt(10)

            def np_():
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                return p
            add_runs(np_(), obj.passage, np_)
        else:
            q, n = obj, extra
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.keep_with_next = True
            r0 = p.add_run(f"Q{n}. ")
            r0.bold = True
            tags = []
            if o["qtype_hint"] and qtype_hint(q):
                tags.append(qtype_hint(q))
            if o["marks"] and q.marks is not None:
                tags.append(fmt_marks(q.marks))
            if tags:
                tr = p.add_run("[" + " · ".join(tags) + "]  ")
                tr.font.size = Pt(9)
                tr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            p = add_runs(p, q.body, lambda: doc.add_paragraph())
            for i, op in enumerate(q.options):
                def opara():
                    pp = doc.add_paragraph()
                    pp.paragraph_format.left_indent = Cm(0.8)
                    pp.paragraph_format.space_after = Pt(1)
                    return pp
                pp = opara()
                lr = pp.add_run(f"({LETTERS[i]}) ")
                lr.bold = True
                add_runs(pp, op.content, opara)
    doc.add_heading("Answer Key", level=2)
    cols = 3 if o["marks"] else 2
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Light Grid Accent 1" if "Light Grid Accent 1" in [s.name for s in doc.styles] else "Table Grid"
    hdr = ["Q", "Answer"] + (["Marks"] if o["marks"] else [])
    for c, h in zip(table.rows[0].cells, hdr):
        c.text = h
    for n, a, mk, _ in m.key:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = str(n), a
        if o["marks"]:
            cells[2].text = fmt_marks(mk).split()[0] if mk is not None else ""
    doc.save(out_path)


# --------------------------------------------------------------------------
# JSON (structured)
# --------------------------------------------------------------------------
def to_json(paper: Paper, m: DocModel) -> dict:
    qs = []
    for kind, obj, extra in m.items:
        if kind != "q":
            continue
        q = obj
        qs.append({
            "number": extra, "original_number": q.number, "question_id": q.qid, "type": q.qtype,
            "marks": q.marks, "text": runs_plain(q.body).strip(), "content": q.body,
            "comprehension_group": q.group,
            "options": [{"label": LETTERS[i], "option_id": op.oid, "text": runs_plain(op.content).strip(),
                         "content": op.content, "correct": op.correct} for i, op in enumerate(q.options)],
            "accepted_answers": q.answers, "answer": answer_str(q),
        })
    groups = {}
    for kind, obj, extra in m.items:
        if kind == "passage":
            groups[obj.gid] = {"questions": list(extra), "text": runs_plain(obj.passage).strip(), "content": obj.passage}
    return {"title": m.title, "source": os.path.basename(paper.source), "paper_meta": paper.meta,
            "section": m.section.name if m.section else None, "subject_check": m.section.subject if m.section else None,
            "comprehension": groups, "questions": qs, "warnings": paper.warnings}


# --------------------------------------------------------------------------
# One-shot conversion
# --------------------------------------------------------------------------
def convert(pdf_path: str, out_root: str, opts: Optional[dict] = None, formats=("md", "html", "docx", "pdf", "json"),
            progress=None) -> dict:
    from .qpaper_engine import extract
    opts = dict(DEFAULT_OPTIONS, **(opts or {}))
    formats = tuple(formats) + ("html",)     # html is always made: it powers the preview
    stem = os.path.splitext(os.path.basename(pdf_path))[0]
    out_dir = os.path.join(out_root, stem)
    if os.path.isdir(out_dir):
        shutil.rmtree(out_dir)
    img_dir = os.path.join(out_dir, "images")
    os.makedirs(out_dir, exist_ok=True)
    paper = extract(pdf_path, img_dir, progress)
    models = build_models(paper, opts)
    produced = []
    for m in models:
        files = {}
        if "md" in formats:
            p = os.path.join(out_dir, m.slug + ".md")
            open(p, "w", encoding="utf-8").write(to_markdown(m, opts))
            files["md"] = p
        if "html" in formats:
            p = os.path.join(out_dir, m.slug + ".html")
            open(p, "w", encoding="utf-8").write(to_html(m, opts))
            files["html"] = p
        if "json" in formats:
            p = os.path.join(out_dir, m.slug + ".json")
            json.dump(to_json(paper, m), open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            files["json"] = p
        if "docx" in formats:
            p = os.path.join(out_dir, m.slug + ".docx")
            to_docx(m, opts, img_dir, p)
            files["docx"] = p
        if "pdf" in formats:
            p = os.path.join(out_dir, m.slug + ".pdf")
            to_pdf(m, opts, img_dir, p)
            files["pdf"] = p
        # markdown bundle (md + only the images it references)
        if "md" in formats:
            zp = os.path.join(out_dir, m.slug + "_markdown.zip")
            md_text = open(files["md"], encoding="utf-8").read()
            used = sorted(set(re.findall(r"images/(img_\d+\.png)", md_text)))
            with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
                z.write(files["md"], f"{m.slug}/{m.slug}.md")
                for u in used:
                    z.write(os.path.join(img_dir, u), f"{m.slug}/images/{u}")
            files["md_zip"] = zp
        n_q = sum(1 for k, *_ in m.items if k == "q")
        produced.append({"title": m.title, "slug": m.slug, "questions": n_q, "files": files})
    sections = [{"index": i, "name": s.name, "subject": s.subject, "partial": s.partial,
                 "questions": len([q for q in s.questions if not q.subject_check])}
                for i, s in enumerate(paper.sections)]
    summary = {"source": os.path.basename(pdf_path), "out_dir": out_dir, "outputs": produced,
               "sections": sections, "warnings": paper.warnings, "meta": paper.meta,
               "guess": guess_title_fields(pdf_path, paper), "options": opts, "profile": "qpaper"}
    json.dump(summary, open(os.path.join(out_dir, "_summary.json"), "w"), indent=1, default=str)
    return summary
